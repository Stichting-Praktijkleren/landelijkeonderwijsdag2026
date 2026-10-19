"""
Bouwt de gegevens voor https://landelijkeonderwijsdag.nl/persoonlijk-programma uit de deelnemers-Excel.

Gebruik:
    python tools/build_persoonlijk_programma.py "C:/pad/naar/2026-10-16 Badges Landelijke Onderwijsdag 2026.xlsx"

Wat het doet
  1. Leest de Excel (kolommen: Voornaam, Achternaam, 'Ronde 1 lange sessie', 'Ronde 1 korte sessie deel 1', ...).
  2. Controleert elke gekozen sessie tegen de programmapagina (bestaat het nummer, klopt de ronde).
  3. Schrijft de DEELNEMERSLIJST VERSLEUTELD (AES-256-GCM) naar netlify/functions/persoonlijk-programma/data.mjs.
     De sleutel staat NIET in de repository: lokaal in .programma-key en op Netlify als geheime omgevingsvariabele.
  4. Schrijft de publieke sessiegegevens (titels, toelichting, sprekers, zalen) naar assets/data/programma-sessies.json.

De Excel zelf wordt nooit in de repository gezet.
"""
import argparse
import base64
import collections
import datetime
import json
import os
import re
import secrets
import sys
import unicodedata
from pathlib import Path

import openpyxl
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

sys.path.insert(0, str(Path(__file__).resolve().parent))
from extract_catalogus import extract  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SLEUTELBESTAND = ROOT / ".programma-key"
SLEUTEL_ENV = "PERSOONLIJK_PROGRAMMA_SLEUTEL"
DATA_UIT = ROOT / "netlify" / "functions" / "persoonlijk-programma" / "data.mjs"
JSON_UIT = ROOT / "assets" / "data" / "programma-sessies.json"
ZALEN_HANDMATIG = ROOT / "tools" / "zalen.json"      # optioneel: {"22": "Zaal Fire 1", ...}
TIJDEN_HANDMATIG = ROOT / "tools" / "tijden.json"    # optioneel: {"r1_kort1": "11.00 – 11.30 uur", ...}

# Dagopbouw: gelijk aan het dagprogramma op programma.html
RONDE_TIJDEN = {1: "11.00 – 12.30 uur", 2: "13.15 – 14.30 uur", 3: "15.00 – 15.30 uur"}
DAG = [
    {"type": "moment", "tijd": "09.00", "titel": "Inloop en ontvangst"},
    {"type": "opening", "tijd": "10.00", "titel": "Opening en keynote",
     "regels": [
         {"label": "Opening", "tekst": "Erwin Pierik, directeur-bestuurder van het Consortium Beroepsonderwijs"},
         {"label": "Keynote", "tekst": "Jolien Dopmeijer: \u2018Het fundament van studentenwelzijn: integraal werken\u2019"},
     ]},
    {"type": "ronde", "ronde": 1},
    {"type": "moment", "tijd": "12.30 \u2013 13.15", "titel": "Lunch", "sub": "Bezoek het Netwerk- en Funplein"},
    {"type": "ronde", "ronde": 2},
    {"type": "moment", "tijd": "14.30 \u2013 15.00", "titel": "Pauze", "sub": "Bezoek het Netwerk- en Funplein"},
    {"type": "ronde", "ronde": 3},
    {"type": "moment", "tijd": "15.30", "titel": "Afsluiting en borrel"},
]

SLOT_RE = re.compile(r"ronde\s*(\d)\s*(lange|korte)\s*sessie(?:\s*deel\s*(\d))?")
ZAAL_RE = re.compile(r"\bzaal\b|zaalnaam|lokaal|ruimte")
CEL_RE = re.compile(r"^\s*(\d+)\s*\.\s*(.*?)\s*(?:[\(\[]\s*zaal\s*[:\-]?\s*([^\)\]]+?)\s*[\)\]])?\s*$", re.S)


def norm(tekst):
    tekst = unicodedata.normalize("NFD", tekst or "")
    tekst = "".join(c for c in tekst if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]", "", tekst)


def slot_id(ronde, soort, deel):
    if soort == "lange":
        return f"r{ronde}_lang"
    return f"r{ronde}_kort{deel}" if deel else f"r{ronde}_kort"


def slot_info(sid):
    m = re.match(r"r(\d)_(lang|kort)(\d?)$", sid)
    ronde, soort, deel = int(m.group(1)), m.group(2), m.group(3)
    if soort == "lang":
        label = "Lange sessie"
    elif deel:
        label = f"Korte sessie \u00b7 deel {deel} van 2"
    else:
        label = "Korte sessie"
    return {"ronde": ronde, "label": label, "tijd": RONDE_TIJDEN[ronde]}


def haal_sleutel():
    waarde = os.environ.get(SLEUTEL_ENV, "").strip()
    if not waarde and SLEUTELBESTAND.exists():
        waarde = SLEUTELBESTAND.read_text(encoding="utf-8").strip()
    if not waarde:
        waarde = secrets.token_hex(32)
        SLEUTELBESTAND.write_text(waarde + "\n", encoding="utf-8")
        print(f"\nNIEUWE SLEUTEL aangemaakt in {SLEUTELBESTAND.name}.")
        print(f"Zet dezelfde waarde als geheime omgevingsvariabele {SLEUTEL_ENV} in Netlify.\n")
    if not re.fullmatch(r"[0-9a-f]{64}", waarde):
        sys.exit("De sleutel moet 64 hexadecimale tekens zijn.")
    return bytes.fromhex(waarde)


def versleutel(gegevens, sleutel):
    iv = secrets.token_bytes(12)
    ciphertext_en_tag = AESGCM(sleutel).encrypt(iv, json.dumps(gegevens, ensure_ascii=False, separators=(",", ":")).encode("utf-8"), None)
    return base64.b64encode(iv + ciphertext_en_tag).decode("ascii")


def lees_excel(pad):
    ws = openpyxl.load_workbook(pad, data_only=True).worksheets[0]
    kop = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    kolommen = {"voornaam": None, "achternaam": None}
    slots, zaalkolommen, onbekende_zaal = {}, {}, []
    for i, naam in enumerate(kop):
        laag = naam.lower()
        if "voornaam" in laag:
            kolommen["voornaam"] = i
        elif "achternaam" in laag:
            kolommen["achternaam"] = i
        m = SLOT_RE.search(laag)
        if m:
            sid = slot_id(int(m.group(1)), m.group(2), m.group(3))
            if ZAAL_RE.search(laag):
                zaalkolommen[sid] = i
            else:
                slots[sid] = i
        elif ZAAL_RE.search(laag):
            onbekende_zaal.append(naam)
    if kolommen["voornaam"] is None or kolommen["achternaam"] is None:
        sys.exit("Kolommen met Voornaam en Achternaam niet gevonden.")
    verwacht = {"r1_lang", "r1_kort1", "r1_kort2", "r2_lang", "r2_kort1", "r2_kort2", "r3_kort"}
    ontbreekt = verwacht - set(slots)
    if ontbreekt:
        sys.exit(f"Sessiekolommen ontbreken in de Excel: {sorted(ontbreekt)}")
    if onbekende_zaal:
        print("LET OP: kolommen met 'zaal' in de kop die niet aan een ronde gekoppeld konden worden:", onbekende_zaal)

    rijen = []
    for r in range(2, ws.max_row + 1):
        waarden = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if all(w in (None, "") for w in waarden):
            continue
        rijen.append((r, waarden))
    return rijen, kolommen, slots, zaalkolommen


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("excel", help="pad naar de deelnemers-Excel")
    ap.add_argument("--data-uit", default=str(DATA_UIT), help="uitvoerbestand voor de versleutelde lijst")
    ap.add_argument("--json-uit", default=str(JSON_UIT), help="uitvoerbestand voor de publieke sessiegegevens")
    args = ap.parse_args()

    excel = Path(args.excel)
    catalogus, typen = extract(with_typen=True)
    rijen, kol, slots, zaalkol = lees_excel(excel)

    deelnemers, zaal_waarnemingen = [], collections.defaultdict(collections.Counter)
    zonder_naam, fouten = 0, []
    for rijnr, w in rijen:
        voornaam = str(w[kol["voornaam"]] or "").strip()
        achternaam = str(w[kol["achternaam"]] or "").strip()
        if not voornaam or not achternaam:
            zonder_naam += 1
            continue
        keuzes = []
        for sid, idx in slots.items():
            cel = w[idx]
            if cel in (None, ""):
                continue
            m = CEL_RE.match(str(cel))
            if not m:
                fouten.append(f"rij {rijnr}: onleesbare sessiecel in {sid}")
                continue
            nr, zaal = int(m.group(1)), (m.group(3) or "").strip()
            if sid in zaalkol and w[zaalkol[sid]] not in (None, ""):
                zaal = str(w[zaalkol[sid]]).strip()
            if nr not in catalogus:
                fouten.append(f"rij {rijnr}: sessie {nr} bestaat niet op programma.html")
                continue
            if catalogus[nr]["ronde"] != slot_info(sid)["ronde"]:
                fouten.append(f"rij {rijnr}: sessie {nr} staat in ronde {catalogus[nr]['ronde']}, niet in {slot_info(sid)['ronde']}")
                continue
            if zaal:
                zaal_waarnemingen[nr][zaal] += 1
            keuzes.append([sid, nr])
        deelnemers.append({"v": voornaam, "a": achternaam, "s": sorted(keuzes)})
    if fouten:
        print("FOUTEN in de Excel:")
        for f in fouten[:40]:
            print("  -", f)
        sys.exit("Afgebroken: los de fouten op en draai opnieuw.")

    # dubbele inschrijvingen van dezelfde persoon samenvoegen (identieke keuzes), twijfelgevallen laten staan
    groepen = collections.defaultdict(list)
    for d in deelnemers:
        groepen[norm(d["v"]) + norm(d["a"])].append(d)
    uniek, samengevoegd, dubbelzinnig = [], 0, 0
    for sleutel, groep in groepen.items():
        keuzesets = {json.dumps(g["s"]) for g in groep}
        if len(keuzesets) == 1:
            uniek.append(groep[0])
            samengevoegd += len(groep) - 1
        else:
            uniek.extend(groep)
            dubbelzinnig += len(groep)

    # zalen: Excel eerst, handmatig bestand vult aan
    zalen = {}
    for nr, teller in zaal_waarnemingen.items():
        if len(teller) > 1:
            print(f"LET OP: sessie {nr} heeft meerdere zalen in de Excel: {dict(teller)}; de meest voorkomende wordt gebruikt.")
        zalen[str(nr)] = teller.most_common(1)[0][0]
    if ZALEN_HANDMATIG.exists():
        for nr, zaal in json.loads(ZALEN_HANDMATIG.read_text(encoding="utf-8")).items():
            zalen.setdefault(str(nr), zaal)

    slots_json = {sid: slot_info(sid) for sid in slots}
    if TIJDEN_HANDMATIG.exists():
        for sid, tijd in json.loads(TIJDEN_HANDMATIG.read_text(encoding="utf-8")).items():
            if sid in slots_json:
                slots_json[sid]["tijd"] = tijd

    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", excel.name)
    stand = f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else datetime.date.fromtimestamp(excel.stat().st_mtime).isoformat()

    publiek = {
        "stand": stand,
        "evenement": {"naam": "Landelijke Onderwijsdag 2026", "datum": "dinsdag 24 november 2026", "locatie": "De Werelt, Lunteren"},
        "dag": DAG,
        "rondes": {str(k): v for k, v in RONDE_TIJDEN.items()},
        "slots": slots_json,
        "typen": typen,
        "zalen": zalen,
        "sessies": {str(nr): s for nr, s in sorted(catalogus.items())},
    }
    Path(args.json_uit).parent.mkdir(parents=True, exist_ok=True)
    Path(args.json_uit).write_text(json.dumps(publiek, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    sleutel = haal_sleutel()
    blob = versleutel({"gemaakt": datetime.datetime.now().isoformat(timespec="seconds"), "deelnemers": uniek}, sleutel)
    Path(args.data_uit).parent.mkdir(parents=True, exist_ok=True)
    Path(args.data_uit).write_text(
        "// Versleutelde deelnemerslijst. Gegenereerd door tools/build_persoonlijk_programma.py. Niet handmatig aanpassen.\n"
        f'export default "{blob}";\n', encoding="utf-8")

    met_keuze = sum(1 for d in uniek if d["s"])
    print(f"Excel: {excel.name} (stand {stand})")
    print(f"Rijen met gegevens: {len(rijen)} | zonder voor- of achternaam (overgeslagen): {zonder_naam}")
    print(f"Deelnemers in de lijst: {len(uniek)} | met minstens 1 sessiekeuze: {met_keuze} | zonder keuze: {len(uniek) - met_keuze}")
    print(f"Dubbele inschrijvingen samengevoegd: {samengevoegd} | namen met afwijkende keuzes (worden als 'meerdere' getoond): {dubbelzinnig}")
    print(f"Zalen bekend voor {len(zalen)} van {len(catalogus)} sessies")
    print(f"Geschreven: {Path(args.data_uit).name} ({len(blob) // 1024} KB versleuteld), {Path(args.json_uit).name}")


if __name__ == "__main__":
    main()
